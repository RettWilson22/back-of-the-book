from __future__ import annotations

import hashlib
from collections.abc import Iterator
from pathlib import Path
from typing import Any

import numpy as np
import pytest
from pydantic import BaseModel

from coursepilot.chunking import Chunk, chunk_pages
from coursepilot.documents import Page
from coursepilot.index import CorpusIndex
from coursepilot.retrieval import Retriever, tokenize


class FakeEmbedder:
    """Deterministic bag-of-words embedding: hashes each token into one of 256 dimensions."""

    model_name = "fake-embedder"

    def encode(self, texts: list[str]) -> np.ndarray:
        vectors = np.zeros((len(texts), 256), dtype=np.float32)
        for row, text in enumerate(texts):
            for token in tokenize(text):
                vectors[row, int(hashlib.md5(token.encode()).hexdigest(), 16) % 256] += 1.0
        norms = np.linalg.norm(vectors, axis=1, keepdims=True)
        return vectors / np.where(norms == 0, 1, norms)


class FakeReranker:
    """Scores a passage by how many query tokens it contains."""

    def __init__(self) -> None:
        self.calls = 0

    def score(self, query: str, texts: list[str]) -> list[float]:
        self.calls += 1
        q = set(tokenize(query))
        return [float(len(q & set(tokenize(t)))) for t in texts]


class FakeLLM:
    name = "fake"

    def __init__(self, reply: str = "", structured: Any = None) -> None:
        self.reply = reply
        self.structured = structured
        self.prompts: list[tuple[str, str]] = []

    def stream_text(self, system: str, user: str) -> Iterator[str]:
        self.prompts.append((system, user))
        words = self.reply.split(" ")
        for i, word in enumerate(words):
            yield word if i == len(words) - 1 else word + " "

    def generate(self, system: str, user: str, schema: type[BaseModel]) -> Any:
        """Return `structured`, or `structured[schema]` when given one response per schema."""
        self.prompts.append((system, user))
        if isinstance(self.structured, dict):
            return self.structured[schema]
        return self.structured


TOY_PAGES = [
    Page("stats.pdf", 1, "The mean is the sum of the values divided by the number of values."),
    Page("stats.pdf", 2, "The median is the middle value when the data are sorted in order."),
    Page("stats.pdf", 3, "Standard deviation measures how spread out values are around the mean."),
    Page("ml.pptx", 1, "A decision tree splits data with a sequence of yes or no questions."),
    Page("ml.pptx", 2, "K-means clustering assigns each point to the nearest of k centroids."),
    Page("ml.pptx", 3, "Overfitting happens when a model memorizes training data noise."),
]


@pytest.fixture
def chunks() -> list[Chunk]:
    return chunk_pages(TOY_PAGES)


@pytest.fixture
def index(chunks: list[Chunk]) -> CorpusIndex:
    return CorpusIndex.build(chunks, FakeEmbedder())


@pytest.fixture
def retriever(index: CorpusIndex) -> Retriever:
    return Retriever(index, FakeEmbedder(), FakeReranker())


def make_pdf(path: Path, pages: list[str]) -> Path:
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path))
    for text in pages:
        y = 750
        for line in text.split("\n"):
            c.drawString(72, y, line)
            y -= 16
        c.showPage()
    c.save()
    return path


def make_pptx(path: Path, slides: list[dict[str, Any]]) -> Path:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for spec in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = spec["title"]
        slide.placeholders[1].text = spec.get("body", "")
        if "table" in spec:
            rows = spec["table"]
            shape = slide.shapes.add_table(
                len(rows), len(rows[0]), Inches(1), Inches(4), Inches(6), Inches(1)
            )
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    shape.table.cell(r, c).text = value
        if "notes" in spec:
            slide.notes_slide.notes_text_frame.text = spec["notes"]
    prs.save(str(path))
    return path
